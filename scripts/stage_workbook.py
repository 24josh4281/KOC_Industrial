"""Read-only mapping review for the supplied management workbook.

Produces LOCAL staging records, never a publishable dataset or official statistics.
No formula evaluation, administrative status inference, or source workbook writes.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'python'))
from kets_dashboard.core import load, ValidationError


def code_version(value):
    match = re.fullmatch(r'([0-9]{2}[A-Z]-[0-9]{3})-Ver([0-9]+)', str(value or '').strip())
    if not match:
        return None, None
    return match.group(1), match.group(2)


def input_date(value):
    if value is None or value in ('', '-'):
        return None
    if isinstance(value, datetime):
        if value.time().isoformat() != '00:00:00':
            raise ValidationError('day-level date has a time component')
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        raise ValidationError('date needs explicit review; no numeric/string guessing')
    return date.fromisoformat(value).isoformat()


def primitive(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def project_rows(sheet, cached, columns, header_row):
    """Retain only selected fields; cell locations and formula/cache provenance survive."""
    from openpyxl.utils.cell import column_index_from_string, get_column_letter
    indices = {label: column_index_from_string(col) - 1 for col, label in columns.items()}
    source = list(sheet.iter_rows(values_only=False))
    values = list(cached.iter_rows(values_only=True))
    for col, label in columns.items():
        if sheet[f'{col}{header_row}'].value != label:
            raise ValidationError(f'{sheet.title}!{col}{header_row}: header mismatch ({label})')
    rows = []
    for row_index in range(header_row, len(source)):
        cells = source[row_index]
        if not any(cells[i].value is not None for i in indices.values()):
            continue
        raw, provenance = {}, {}
        for label, i in indices.items():
            cell = cells[i]
            value = values[row_index][i] if row_index < len(values) else None
            raw[label] = primitive(value)
            provenance[label] = {'cell': f'{get_column_letter(i + 1)}{row_index + 1}', 'formula': cell.value if cell.data_type == 'f' else None,
                                 'raw_value': primitive(cell.value), 'cached_value': primitive(value),
                                 'cell_type': cell.data_type}
        rows.append({'source_sheet': sheet.title, 'source_row': row_index + 1,
                     'raw': raw, 'cells': provenance})
    return rows


def map_records(reference, applications, internal, *, observed_on, source_hash):
    """Pure mapping: candidates remain separate from the dashboard dataset contract."""
    date.fromisoformat(observed_on)
    rules = load(ROOT / 'config/workbook_mapping.json')
    issues, methods, versions, app_candidates, histories = [], [], [], [], []
    methods_by_code = {}
    for record in reference:
        raw = record['raw']
        code, version = code_version(raw['방법론 고유번호'])
        row = record['source_row']
        if not code:
            issues.append({'sheet': record['source_sheet'], 'row': row, 'issue': 'unparsed_methodology_code'})
            continue
        mid = 'M-' + code
        scope = rules['scope_candidates'].get(raw['산업/발전 구분'], 'unclassified')
        candidate = {'id': mid, 'code': code, 'name': raw['방법론명'], 'scope': scope,
                     'scope_basis': f"{record['source_sheet']}!J{row}: {raw['산업/발전 구분']} (내부 분류·승인 대기)",
                     'official_sector': raw['대분류'], 'technology': '',
                     'first_registered_on': None, 'registration_confirmed_on': None,
                     'recorded_on': observed_on, 'source_id': 'SOURCE-INPUT', 'public_allowed': False}
        if code not in methods_by_code:
            methods_by_code[code] = candidate
            methods.append(candidate)
        else:
            previous = methods_by_code[code]
            if any(previous[k] != candidate[k] for k in ('name', 'scope', 'official_sector')):
                issues.append({'sheet': record['source_sheet'], 'row': row, 'issue': 'methodology_metadata_conflict', 'code': code})
        try:
            effective = input_date(raw['방법론유효시작일'])
        except ValidationError:
            effective = None
            issues.append({'sheet': record['source_sheet'], 'row': row, 'issue': 'invalid_effective_date'})
        versions.append({'id': f'V-{code}-{version}', 'methodology_id': mid, 'version': version,
                         'change_type': rules['type_candidates'].get(raw['신규/개정']), 'registered_on': None,
                         'effective_from': effective, 'recorded_on': observed_on, 'source_id': 'SOURCE-INPUT'})
    # Source row number is a proposed local key, not an official ORS application ID.
    internal_by_number = {}
    internal_by_key = {}
    for record in internal:
        number = record['raw']['번호']
        if number in internal_by_number:
            raise ValidationError('duplicate internal reference number')
        internal_by_number[number] = record
        key = (record['raw']['방법론명'], record['raw']['최초신청일'])
        internal_by_key.setdefault(key, []).append(record)
    numbers = set()
    for record in applications:
        raw, row = record['raw'], record['source_row']
        number = raw['번호']
        if number is None or number in numbers:
            raise ValidationError('missing/duplicate ORS reference number')
        numbers.add(number)
        proposed_id = f'LOCAL-ORS-METHOD-{number}'
        code, _ = code_version(raw['방법론고유번호'])
        method = methods_by_code.get(code)
        matches = internal_by_key.get((raw['방법론명'], raw['최초신청일']), [])
        related = matches[0] if len(matches) == 1 else None
        blockers = []
        if related is None:
            blockers.append('internal_join_not_verified')
        if related and related['raw']['신청유형'] != raw['신청유형']:
            blockers.append('application_type_conflict')
        kind = rules['type_candidates'].get(raw['신청유형'])
        if kind is None:
            blockers.append('unknown_application_type')
        if kind == 'revision' and method is None:
            blockers.append('revision_methodology_link_missing')
        try:
            submitted = input_date(raw['최초신청일'])
        except ValidationError:
            submitted = None
        if submitted is None or submitted > observed_on:
            blockers.append('submission_date_requires_review')
        candidate = {'id': proposed_id, 'methodology_id': method['id'] if method else None,
                     'title': raw['방법론명'], 'type': kind,
                     'scope': method['scope'] if method else 'unclassified', 'submitted_on': submitted,
                     'recorded_on': observed_on, 'owner': '', 'source_id': 'SOURCE-INPUT', 'public_allowed': False}
        app_candidates.append({'proposed': candidate, 'source_row': row,
                               'internal_source_row': related['source_row'] if related else None,
                               'join_rule': 'unique exact title AND first-submission date; not source row number',
                               'official_status_text': raw['처리상태'],
                               'internal_status_text': related['raw']['진행상태'] if related else None,
                               'internal_type': related['raw']['신청유형'] if related else None,
                               'blockers': blockers})
        for blocker in blockers:
            issues.append({'sheet': record['source_sheet'], 'row': row, 'issue': blocker})
        for line_no, line in enumerate(str(raw.get('이력') or '').splitlines(), 1):
            if not line.strip():
                continue
            match = re.fullmatch(r'(\d{4}-\d{2}-\d{2}) (\d{2}:\d{2}(?::\d{2})?) - (.+)', line.strip())
            # Original timestamp and row order retained. No invented stage or transition.
            histories.append({'application_id_candidate': proposed_id, 'source_row': row,
                              'line': line_no, 'raw_text': line,
                              'occurred_on_candidate': match.group(1) if match else None,
                              'time_raw': match.group(2) if match else None,
                              'official_status_text': match.group(3) if match else None,
                              'stage': None, 'action': None, 'recorded_on': observed_on,
                              'review_required': True})
    duplicates = [key for key, count in Counter((v['methodology_id'], v['version']) for v in versions).items() if count > 1]
    for key in duplicates:
        issues.append({'issue': 'duplicate_methodology_version', 'key': list(key)})
    return {
        'staging_version': '1.0.0', 'status': 'mapping_review_required',
        'observed_on': observed_on, 'source_sha256': source_hash,
        'notice': '원장 매핑 검토용. 공식 통계·확정 dataset 아님. 등록일·적용상태·신청단계 미확정.',
        'proposed_methodologies': methods, 'proposed_versions': versions,
        'proposed_applications': app_candidates, 'raw_review_history': histories,
        'proposed_status_events': [],
        'issues': issues,
        'summary': {'reference_rows': len(reference), 'methodology_candidates': len(methods),
                    'version_candidates': len(versions), 'application_rows': len(applications),
                    'history_lines': len(histories), 'issue_counts': dict(Counter(x['issue'] for x in issues))},
        'promotion_blockers': ['등록사실 확인일·최초등록일 근거', '적용상태 근거',
                               '개정 신청의 방법론 ID 연결', '신청유형 불일치 확인',
                               '공식 상태와 내부 상태의 단계 매핑 승인', '미분류 신청의 산업/전환 범위',
                               '기록일은 수입일이며 원래 입력일이 아님', '로컬 신청 ID 안정성 승인'],
        'source_records': {'reference': reference, 'applications': applications, 'internal': internal},
    }


def stage_workbook(path, observed_on):
    import openpyxl
    from openpyxl.utils.cell import get_column_letter
    config = load(ROOT / 'config/workbook_mapping.json')
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ref_columns = {get_column_letter(i + 1): label for i, label in enumerate(config['reference_headers'])}
        records = []
        for prefix, columns in [('reference', ref_columns), ('application', config['application_columns']), ('internal', config['internal_columns'])]:
            name = config[prefix + '_sheet']
            if name not in book.sheetnames:
                raise ValidationError(f'missing sheet: {name}')
            records.append(project_rows(book[name], cached[name], columns, config[prefix + '_header_row']))
        result = map_records(*records, observed_on=observed_on, source_hash=sha256(Path(path).read_bytes()).hexdigest())
        result['source_filename'] = Path(path).name
        result['proposed_sources'] = [{'id': 'SOURCE-INPUT', 'label': Path(path).name,
                                      'url': None, 'kind': 'manual', 'access': 'internal',
                                      'retrieved_on': observed_on, 'file_sha256': result['source_sha256']}]
        result['excluded_sheets'] = [s for s in book.sheetnames if s not in {config[x + '_sheet'] for x in ('reference', 'application', 'internal')}]
        return result
    finally:
        book.close()
        cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('--observed-on', required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    try:
        output = args.out.resolve()
        private_roots = [ROOT / 'local', ROOT / 'data' / 'staging', ROOT / 'data' / 'private']
        if not any(output.is_relative_to(p.resolve()) for p in private_roots):
            raise ValidationError('actual data outputs must stay in local/ or data/staging/ or data/private/')
        if output.exists():
            raise ValidationError('output exists; choose a new _v2/_review filename')
        result = stage_workbook(args.input, args.observed_on)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(json.dumps({'status': result['status'], 'summary': result['summary']}, ensure_ascii=False, indent=2))
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(2, f'ERROR: {exc}\n')


if __name__ == '__main__':
    main()
